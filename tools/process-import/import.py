"""Generate init/05_process_recipes_seed.sql from the "Processing Flips" sheet.

The sheet (a friend's Google Sheet, exported to .xlsx) encodes each processing
recipe as one hand-written profit formula on its Main tab, e.g.

    C398  =3*roundup(X397*0.98,0)-Y398        (BCP -> 3 bandosian components)
    E398  =C398*60/1000000                    (60 actions/hr -> m gp/hr)

where X/Y are the high/low prices of the item named in column W of that row.
This script turns every formula into structured legs (item_id, qty), a per-
action fee and an actions/hr figure, then proves each parse by re-evaluating
the legs at the sheet's own cached prices and comparing to the sheet's cached
profit. Anything that doesn't reproduce is held back, never guessed.

Usage (no system python on eldo; run from ge-data/):

    nix-shell -p 'python3.withPackages(p:[p.openpyxl p.sympy])' --run \\
      "python -I tools/process-import/import.py 'Processing Flips.xlsx' mapping.json"

mapping.json is the OSRS Wiki /mapping (same source as the `items` table):
    curl -A '<contact>' https://prices.runescape.wiki/api/v1/osrs/mapping -o mapping.json
"""
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
import sympy as sp
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parents[2]
RELATIONS_SEED = ROOT / "init" / "03_item_relations_seed.sql"
OUT_SQL = ROOT / "init" / "05_process_recipes_seed.sql"

# ── Curation ────────────────────────────────────────────────────────────────
# Sheet rows whose formula picks the cheapest of several sources (masori/torva
# take plates/components from whichever armour piece is cheapest) can't be
# expressed as fixed legs, but item_relations already models them with the
# plate/component as its own tradeable item. Link those by name instead.
# (Linking also keeps out a sheet bug: it charges 3 bandosian components for
# the Torva platebody; the Wiki and item_relations say 2.)
LINK = {
    322: "Accursed sceptre (u)",
    405: "Masori mask (f)",
    406: "Masori body (f)",
    407: "Masori chaps (f)",
    409: "Torva full helm (restore)",
    410: "Torva platebody (restore)",
    411: "Torva platelegs (restore)",
    # Sold untaxed with a hand-written "-5000000": that's the GE tax cap, not a fee.
    413: "Zaryte crossbow",
}
# Accepted although the sheet's cached value disagrees, because the sheet is wrong.
ACCEPT = {
    56: "sheet formula has min(,$Y$30), which drops the 2x rune cost; this seed includes it",
}
DEFER = {
    "Crystal Stuffs": "values crystal shards through a derived price; shards are untradeable, needs a valuation model",
}
DEFER_ROWS = {
    109: "high-alch output (coins), not a GE sale",
    111: "high-alch output (coins), not a GE sale",
    113: "high-alch output (coins), not a GE sale",
    115: "high-alch output (coins), not a GE sale",
}
RENAME = {
    "Making Aram from scratch": "Making Arma from scratch",
    "Chcolate Dust": "Chocolate Dust",
    "Rune Incindiary Cannonballs": "Rune Incendiary Cannonballs",
    "Swampbark Set from Stractch": "Swampbark Set from Scratch",
    "Repair Torva Platbody (3)": "Repair Torva Platebody (3)",
}
RECATEGORY = {"Indendiary Cannonballs": "Incendiary Cannonballs"}
# The pyre rows sit under the sheet's GWD header with no header of their own.
RECATEGORY_ROWS = {n: "Pyre Logs" for n in range(421, 434)}


def rows(a, b):
    return range(a, b + 1)


# Wiki-checked corrections to legs the sheet gets wrong. The self-check only
# proves a parse matches the SHEET; these make it match the GAME. Applied after
# parsing, before linking. Keys are sheet rows; each fix says why in `note`.
#   in_set / out_set: {item_id: qty} (adds the leg if missing)
#   replace:          {old_item_id: new_item_id} on either side (must exist)
TAN = dict(in_set={9075: 2, 554: 5}, note="Tan Leather costs 2 astral + 1 nature + 5 fire per cast (Wiki); the sheet charges 1 astral and no fire.")
DEGRIME = dict(in_set={561: 2, 557: 4}, note="Degrime costs 2 nature + 4 earth per cast (Wiki); the sheet charges 1 nature and no earth.")
CHEM = dict(in_set={21163: 0.015}, note="Amulet of chemistry: 15% proc per potion over 10 charges = 0.015 amulets/action; the sheet omits the /10 here.")
CLEAN_CADANTINE = dict(replace={215: 265}, note="Zahur makes unfinished potions from clean herbs only (Wiki); the sheet buys grimy cadantine.")
ZENYTE = dict(in_set={2357: 1, 564: 1}, note="Adds the gold bar and the Lvl-7 Enchant cosmic rune the sheet leaves out.")
FISH_CRATE = dict(replace={32366: 32364}, note="Emptying a filled crate returns a camphor crate (Wiki); the sheet values it as an empty fish crate.")
LEG_FIX = {
    **{n: TAN for n in rows(24, 27)},
    **{n: DEGRIME for n in rows(60, 74)},
    200: CHEM,  # Making Bastion Potion
    202: CHEM,  # Making Battlemage Potion
    # The "+ unf" rows also chain Cadantine Blood unf, so they carry both errors.
    **{n: dict(in_set=CHEM["in_set"], replace=CLEAN_CADANTINE["replace"],
               note=CHEM["note"] + " " + CLEAN_CADANTINE["note"]) for n in (201, 203)},
    246: CLEAN_CADANTINE,
    **{n: ZENYTE for n in rows(303, 310)},
    **{n: dict(in_set={**ZENYTE["in_set"], 1759: 1}, note=ZENYTE["note"] + " Torture is an amulet, so it also needs a ball of wool.")
       for n in (305, 306)},
    **{n: FISH_CRATE for n in rows(164, 169)},
    221: dict(out_set={11994: 4.5}, note="A lava scale crushes to 3-6 shards; 4.5 is the mean (the sheet uses the minimum, 3)."),
}
# Sheet rows whose "actions/hr" is really a buy limit (e.g. 13000*12, 200*12)
# or an NPC bulk exchange, not player speed. They're bounded by buy limits
# and market volume instead.
VOLUME_BOUND_ROWS = {*rows(76, 90), 299, 300, *rows(36, 43), 172, 216}
VOLUME_BOUND_NOTE = "The sheet's actions/hr here came from a buy limit or an NPC bulk exchange, not player speed; bounded by buy limits and market volume instead."
NOTES = {
    **{n: "The sheet charges 1 astral per sapling; Humidify waters a whole inventory per cast, so that's an overestimate. The filled plant pot (~3 gp) is left out."
       for n in rows(36, 43)},
    99: "Same legs as Cooking Karambwans at 1-tick speed.",
}


# ── Sheet access ────────────────────────────────────────────────────────────
def grid(ws, cols):
    g = {}
    for n, row in enumerate(ws.iter_rows(max_col=30), start=1):
        for i, c in enumerate(row):
            col = get_column_letter(i + 1)
            if col in cols and c.value is not None:
                v = c.value
                g[(col, n)] = v.text if isinstance(v, ArrayFormula) else v
    return g


def load(path):
    wbf = openpyxl.load_workbook(path, read_only=True, data_only=False)
    wbv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cols = set("ABCDEIKTUWXY")
    formulas, values = grid(wbf["Main"], cols), grid(wbv["Main"], cols)
    # Item Data: G id, H name, J buy price (high), L sell price (low)
    idv = grid(wbv["Item Data"], set("GHJL"))
    items = {}
    for (col, n), v in idv.items():
        gid = idv.get(("G", n))
        if col == "H" and isinstance(gid, (int, float)):
            items[str(v).lower()] = dict(id=int(gid), name=v, J=idv.get(("J", n)), L=idv.get(("L", n)))
    return formulas, values, items


# ── Formula translation ─────────────────────────────────────────────────────
def translate(f):
    """Excel formula -> python expression over REF/RNG/MAX/MIN/ROUNDUP/XL."""
    s = f[1:] if f.startswith("=") else f
    s = re.sub(r"_xlfn\.xlookup\(", "XL(", s, flags=re.I)
    s = re.sub(r"'Item Data'!([A-Z]+):[A-Z]+", r"'\1'", s)
    s = re.sub(r"\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)", r"RNG('\1',\2,'\3',\4)", s)
    s = re.sub(r"(?<![A-Za-z_'])\$?([A-Z]{1,2})\$?(\d+)(?![\d(])", r"REF('\1',\2)", s)
    s = re.sub(r"\b(roundup|max|min|if)\(", lambda m: m.group(1).upper() + "(", s, flags=re.I)
    return s.replace("(,", "(None,")


SAFE_NAMES = {"REF", "RNG", "MAX", "MIN", "ROUNDUP", "XL", "TAXV", "None"}


def safe_eval(s, env):
    """eval() a translated formula, which comes from an untrusted sheet: only
    our helper names, numbers, arithmetic and quoted strings may appear, and
    builtins are off."""
    bare = re.sub(r"'[^']*'|\"[^\"]*\"", "", s)
    names = set(re.findall(r"[A-Za-z_][A-Za-z_0-9]*", bare))
    rest = re.sub(r"[A-Za-z_][A-Za-z_0-9]*", "", bare)
    if names - SAFE_NAMES or set(rest) - set("0123456789.+-*/(), "):
        raise ValueError(f"unsafe formula: {s[:80]!r}")
    return eval(s, {"__builtins__": {}, **env})


def flat(a):
    for x in a:
        if isinstance(x, list):
            yield from flat(x)
        else:
            yield x


def numeric_env(values, items):
    """Evaluate a formula exactly as Sheets does, at the cached cell values."""
    def REF(c, n):
        return values.get((c, n))

    def RNG(c1, n1, c2, n2):
        cs = [get_column_letter(i) for i in range(column_index_from_string(c1), column_index_from_string(c2) + 1)]
        return [values.get((c, n)) for c in cs for n in range(n1, n2 + 1)]

    def nums(a):
        return [x if x is not None else 0 for x in flat(a)]  # a blank argument counts as 0

    return dict(
        REF=REF, RNG=RNG,
        MAX=lambda *a: max(nums(a)), MIN=lambda *a: min(nums(a)),
        ROUNDUP=lambda x, d: math.ceil(x) if x >= 0 else -math.ceil(-x),
        XL=lambda name, k, r: items[name.lower()][r],
    )


TAX = sp.Symbol("TAX")


def symbolic(f, formulas, values, items, flags, depth=0):
    """Formula -> sympy expression: prices become H_<row>/L_<row>/M_<row>
    symbols (high, low, min of the two), XLOOKUPs become H_xl<id>/L_xl<id>,
    and the literal 0.98 becomes TAX so taxed outputs stay distinguishable."""
    if depth > 3:
        raise ValueError("composite nesting too deep")

    def REF(c, n):
        if c == "X":
            return sp.Symbol(f"H_{n}")
        if c == "Y":
            return sp.Symbol(f"L_{n}")
        if c == "C":
            flags.add("composite")
            return symbolic(formulas[("C", n)], formulas, values, items, flags, depth + 1)
        flags.add(f"ref_{c}")
        return sp.Float(values.get((c, n)) or 0)

    def RNG(c1, n1, c2, n2):
        if (c1, c2) == ("X", "Y") and n1 == n2:
            return [sp.Symbol(f"M_{n1}")]
        flags.add("range_alt")
        return [REF(c1, n) for n in range(n1, n2 + 1)]

    def pick(a):
        a = [x for x in flat(a) if x is not None]
        if len(a) != 1:
            flags.add("alt_inputs")
        return a[0]

    def XL(name, k, r):
        return sp.Symbol(f"{'H' if r == 'J' else 'L'}_xl{items[name.lower()]['id']}")

    s = re.sub(r"(?<![\d.])0\.98(?!\d)", "TAXV", translate(f))
    env = dict(REF=REF, RNG=RNG, MAX=lambda *a: pick(a), MIN=lambda *a: pick(a),
               ROUNDUP=lambda x, d: x, XL=XL, TAXV=TAX)
    return sp.sympify(safe_eval(s, env))


def resolve(sym, values, items, by_id):
    """Symbol -> (item_id, sheet price used for the self-check)."""
    kind, key = sym.name.split("_", 1)
    if key.startswith("xl"):
        it = by_id[int(key[2:])]
        return it["id"], it["J"] if kind == "H" else it["L"]
    n = int(key)
    it = items.get(str(values.get(("W", n))).lower())
    hi, lo = values.get(("X", n)), values.get(("Y", n))
    price = {"H": hi, "L": lo, "M": min(p for p in (hi, lo) if p is not None)}[kind]
    return (it["id"] if it else None), price


def to_legs(expr, values, items, by_id, flags):
    """Read legs off the expanded expression. Returns (outs, ins, fee, value):
    outs/ins as {item_id: qty}, fee per action, and `value` = those legs and
    fee re-priced at the sheet's cached prices (outputs at 0.98 x high), the
    figure the caller compares to the sheet. Built term by term from what was
    extracted, so a term the extraction drops shows up as a mismatch."""
    e = sp.expand(expr)
    outs, ins = {}, {}
    value = 0.0
    # The constant must not depend on TAX: a fee inside a taxed bracket,
    # e.g. ROUNDUP((X5-1000)*0.98), would otherwise vanish from fee_gp.
    const = sp.Poly(e.subs({s: 0 for s in e.free_symbols - {TAX}}), TAX)
    if const.degree() > 0:
        flags.add("taxed_constant")
    fee = -float(const.coeff_monomial(1))
    value -= fee
    for s in sorted(e.free_symbols - {TAX}, key=lambda s: s.name):
        poly = sp.Poly(e.coeff(s), TAX)
        if poly.degree() > 1:
            flags.add("nonlinear")
        taxed, untaxed = float(poly.coeff_monomial(TAX)), float(poly.coeff_monomial(1))
        iid, price = resolve(s, values, items, by_id)
        if iid is None:
            flags.add("unmapped_item")
        if taxed > 0:
            outs[iid] = outs.get(iid, 0) + taxed
            value += taxed * 0.98 * price
        if untaxed > 0:
            flags.add("untaxed_output")
        if untaxed < 0:
            ins[iid] = ins.get(iid, 0) - untaxed
            value += untaxed * price
        if taxed < 0:
            flags.add("taxed_input")
        if s.name[0] == "L" and taxed > 0:
            flags.add("sells_at_low")
    return outs, ins, fee, value


def net(outs, ins):
    """Cancel an item that is both sold and bought in one action. The sheet
    chains steps by adding their profits (C366+C363), which sells the
    intermediate at high and buys it back at low: a spread nobody trades."""
    netted = []
    for iid in set(outs) & set(ins):
        n = min(outs[iid], ins[iid])
        outs[iid] -= n
        ins[iid] -= n
        netted.append(iid)
    outs = {k: v for k, v in outs.items() if v > 1e-9}
    ins = {k: v for k, v in ins.items() if v > 1e-9}
    return outs, ins, netted


def apply_fix(outs, ins, fix):
    for old, new in fix.get("replace", {}).items():
        side = outs if old in outs else ins if old in ins else None
        if side is None:
            raise ValueError(f"replace: item {old} not in legs")
        side[new] = side.pop(old)
    for iid, qty in fix.get("in_set", {}).items():
        ins[iid] = qty
    for iid, qty in fix.get("out_set", {}).items():
        outs[iid] = qty


def leg_list(d):
    return [{"item_id": k, "qty": round(v, 6)} for k, v in sorted(d.items())]


def throughput(f):
    """Column E -> (actions_per_hour, throughput kind)."""
    if not isinstance(f, str):
        return None, "unknown"
    s = translate(f)
    product = lambda p: math.prod(float(x) for x in p.split("*"))
    m = re.fullmatch(r"REF\('C',\d+\)\*([\d.]+(?:\*[\d.]+)*)/(\d+)(\*12)?", s)
    if m:
        return round(product(m.group(1)) * 1_000_000 / int(m.group(2)) * (12 if m.group(3) else 1), 3), "fixed"
    m = re.fullmatch(r"([\d.]+(?:\*[\d.]+)*)\*REF\('C',\d+\)/(\d+)", s)
    if m:
        return round(product(m.group(1)) * 1_000_000 / int(m.group(2)), 3), "fixed"
    m = re.search(r"(\d+)\*(\d+)\)", s)
    if m and "'T'" in s:
        return int(m.group(1)) * int(m.group(2)), "volume_capped"
    if "'T'" in s:
        return None, "volume_bound"
    return None, "unknown"


# ── item_relations (for linking recipes we already curate) ──────────────────
def existing_relations():
    seed = RELATIONS_SEED.read_text()
    rels = {}
    for m in re.finditer(r"\('(\w+)',\s*'((?:[^']|'')*)',\s*(true|false),\s*'(\[.*?\])',\s*'(\[.*?\])'", seed):
        key = lambda legs: tuple(sorted((l["item_id"], l["qty"]) for l in json.loads(legs)))
        rels[m.group(2).replace("''", "'")] = dict(reversible=m.group(3) == "true",
                                                   ins=key(m.group(4)), outs=key(m.group(5)))
    return rels


def find_link(outs, ins, rels):
    k = lambda legs: tuple(sorted((l["item_id"], l["qty"]) for l in legs))
    ko, ki = k(outs), k(ins)
    for name, r in rels.items():
        if (r["ins"], r["outs"]) == (ki, ko):
            return name, "forward"
        if r["reversible"] and (r["outs"], r["ins"]) == (ki, ko):
            return name, "reverse"
    return None


# ── Main ────────────────────────────────────────────────────────────────────
def main(xlsx, mapping_path):
    formulas, values, items = load(xlsx)
    by_id = {v["id"]: v for v in items.values()}
    mapping = {m["id"]: m for m in json.load(open(mapping_path))}
    rels = existing_relations()

    recipes, held = [], []
    category = ""
    for n in range(4, 500):
        a = values.get(("A", n))
        if a and str(a).strip():
            category = str(a).strip()
        name, f = values.get(("B", n)), formulas.get(("C", n))
        if not name:
            continue
        if not isinstance(f, str) or not f.startswith("=") or re.search(r"\bif\(", f, re.I):
            continue  # section headers, helper lookups
        if re.fullmatch(r"=-?(1\*)?\$?[DY]\$?\d+", f):
            held.append((n, name, "mirror of another row (an 'undo set' or a helper price), not a recipe"))
            continue
        cat = RECATEGORY_ROWS.get(n) or RECATEGORY.get(category, category)
        if category in DEFER or n in DEFER_ROWS:
            held.append((n, name, DEFER.get(category) or DEFER_ROWS[n]))
            continue
        aph, kind = throughput(formulas.get(("E", n)))
        notes = [NOTES[n]] if n in NOTES else []
        if n in VOLUME_BOUND_ROWS:
            aph, kind = None, "volume_bound"
            notes.append(VOLUME_BOUND_NOTE)
        rec = dict(row=n, name=RENAME.get(name, name), category=cat, aph=aph, throughput=kind,
                   relation=None, direction=None, inputs=None, outputs=None, fee=0.0, notes=None)

        if n in LINK:
            if LINK[n] not in rels:
                sys.exit(f"row {n}: LINK target {LINK[n]!r} not in item_relations seed")
            rec["relation"], rec["direction"] = LINK[n], "forward"
            rec["notes"] = " ".join(notes) or None
            recipes.append(rec)
            continue

        flags = set()
        expr = symbolic(f, formulas, values, items, flags)
        outs, ins, fee, value = to_legs(expr, values, items, by_id, flags)
        sheet_c = values.get(("C", n))
        exact = safe_eval(translate(f), numeric_env(values, items))
        tol = sum(outs.values()) + 1  # ROUNDUP slack: up to 1 gp per unit sold
        if abs(exact - sheet_c) > 1e-6:
            flags.add("translation_mismatch")
        if abs(value - sheet_c) > tol and n not in ACCEPT:
            flags.add("legs_mismatch")
        for iid in [*outs, *ins]:
            if iid not in mapping:
                flags.add("item_not_in_mapping")
        hard = flags - {"composite"}
        if hard:
            held.append((n, name, "parse: " + ",".join(sorted(hard))))
            continue

        # Past the self-check, the legs match the sheet. Now make them match
        # the game: cancel chained intermediates, then apply Wiki corrections.
        outs, ins, netted = net(outs, ins)
        if netted:
            notes.insert(0, "Chains two steps; the intermediate (item "
                         + ", ".join(str(i) for i in sorted(netted)) + ") is made and used, not traded.")
        if n in LEG_FIX:
            apply_fix(outs, ins, LEG_FIX[n])
            notes.append(LEG_FIX[n]["note"])
        if n in ACCEPT:
            notes.append(ACCEPT[n])
        if not outs or not ins:
            sys.exit(f"row {n} {name}: no {'outputs' if not outs else 'inputs'} left after netting/fixes")
        for iid in [*outs, *ins]:
            if iid not in mapping:
                sys.exit(f"row {n} {name}: corrected leg item {iid} not in mapping")
        rec["fee"] = round(fee, 4)
        rec["notes"] = " ".join(notes) or None
        outs_l, ins_l = leg_list(outs), leg_list(ins)
        link = find_link(outs_l, ins_l, rels)
        if link:
            rec["relation"], rec["direction"] = link
        else:
            rec["inputs"], rec["outputs"] = ins_l, outs_l
        recipes.append(rec)

    names = Counter(r["name"] for r in recipes)
    dup = [k for k, c in names.items() if c > 1]
    if dup:
        sys.exit(f"duplicate recipe names: {dup}")
    write_sql(recipes, held)
    print(f"{len(recipes)} recipes written ({sum(1 for r in recipes if r['relation'])} linked to item_relations), "
          f"{len(held)} held back", file=sys.stderr)
    for n, name, why in held:
        print(f"  held row {n} {name}: {why}", file=sys.stderr)
    print(Counter(r["throughput"] for r in recipes), file=sys.stderr)


def q(s):
    return "NULL" if s is None else "'" + str(s).replace("'", "''") + "'"


def num(x):
    if x is None:
        return "NULL"
    return str(int(x)) if float(x).is_integer() else repr(round(float(x), 6))


def legs_sql(legs):
    if legs is None:
        return "NULL"
    return q(json.dumps([{"item_id": l["item_id"], "qty": float(l["qty"]) if not float(l["qty"]).is_integer() else int(l["qty"])}
                         for l in legs], separators=(",", ":")))


def write_sql(recipes, held):
    lines = [
        "-- process_recipes seed: GENERATED by tools/process-import/import.py from the",
        "-- \"Processing Flips\" sheet (Main tab). Don't hand-edit rows; change the",
        "-- curation tables in the importer and re-run it. Each unlinked row's parsed",
        "-- legs reproduced the sheet's own cached profit/action (except rows listed in",
        "-- ACCEPT), and were then corrected where the sheet disagrees with the game",
        "-- (chained intermediates cancelled, Wiki fixes in LEG_FIX); every correction",
        "-- is described in the row's notes.",
        "--",
        "-- Re-runnable and exact: one transaction upserts every row and deletes",
        "-- importer rows (source 'Main!...') whose name is no longer generated.",
        "",
        "BEGIN;",
        "",
        "INSERT INTO process_recipes (name, category, relation_name, relation_direction, inputs, outputs,",
        "                             fee_gp, actions_per_hour, throughput, source, notes) VALUES",
    ]
    vals, last = [], None
    for r in recipes:
        if r["category"] != last:
            vals.append(f"\n-- ── {r['category']} ──")
            last = r["category"]
        vals.append(f"({q(r['name'])}, {q(r['category'])}, {q(r['relation'])}, {q(r['direction'])}, {legs_sql(r['inputs'])}, "
                    f"{legs_sql(r['outputs'])}, {num(r['fee'])}, {num(r['aph'])}, {q(r['throughput'])}, "
                    f"'Main!B{r['row']}', {q(r['notes'])}),")
    # drop the trailing comma on the last VALUES row
    for i in range(len(vals) - 1, -1, -1):
        if vals[i].endswith("),"):
            vals[i] = vals[i][:-1]
            break
    lines += vals
    lines += [
        "ON CONFLICT (name) DO UPDATE SET",
        "  category = EXCLUDED.category, relation_name = EXCLUDED.relation_name,",
        "  relation_direction = EXCLUDED.relation_direction,",
        "  inputs = EXCLUDED.inputs, outputs = EXCLUDED.outputs, fee_gp = EXCLUDED.fee_gp,",
        "  actions_per_hour = EXCLUDED.actions_per_hour, throughput = EXCLUDED.throughput,",
        "  source = EXCLUDED.source, notes = EXCLUDED.notes;",
        "",
        "DELETE FROM process_recipes",
        "WHERE source LIKE 'Main!%'",
        "  AND name NOT IN (" + ", ".join(q(r["name"]) for r in recipes) + ");",
        "",
        "COMMIT;",
        "",
        "-- Held back (not imported):",
    ]
    lines += [f"--   Main!B{n} {name}: {why}" for n, name, why in held]
    OUT_SQL.write_text("\n".join(lines) + "\n")


if __name__ == "__main__":
    main(*sys.argv[1:3])
